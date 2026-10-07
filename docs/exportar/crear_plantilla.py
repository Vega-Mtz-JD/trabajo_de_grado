"""Genera plantilla_upea.docx (reference-doc de pandoc) con el formato del Reglamento UPEA 2025.

Art. 34/35: Carta, Arial 11 (10 en tablas/figuras/notas), interlineado 2,
espacio entre párrafos 6–12 pt, márgenes sup/inf/der 2.55 cm e izq 3 cm,
numeración arriba a la derecha (APA). Requiere pandoc y python-docx.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

DESTINO = Path(__file__).with_name("plantilla_upea.docx")


def campo_pagina(parrafo):
    """Inserta el campo PAGE (número de página) en un párrafo."""
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def quitar_tema(rfonts):
    # Los estilos de pandoc usan la fuente del tema (asciiTheme), que tiene prioridad sobre "Arial".
    for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(atributo), None)


def arial(font, rpr):
    font.name = "Arial"
    rpr.rFonts.set(qn("w:eastAsia"), "Arial")
    quitar_tema(rpr.rFonts)


def fuente(estilo, tam, negrita=None):
    arial(estilo.font, estilo.element.get_or_add_rPr())
    estilo.font.size = Pt(tam)
    estilo.font.color.rgb = None
    if negrita is not None:
        estilo.font.bold = negrita


def main():
    base = subprocess.run(
        ["pandoc", "--print-default-data-file", "reference.docx"],
        check=True, capture_output=True,
    ).stdout
    tmp = DESTINO.with_suffix(".tmp.docx")
    tmp.write_bytes(base)
    doc = Document(tmp)

    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.top_margin = sec.bottom_margin = sec.right_margin = Cm(2.55)
        sec.left_margin = Cm(3)
        cab = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
        cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        campo_pagina(cab)

    # Búsqueda directa por nombre: python-docx traduce "Heading 1" a "heading 1" y no
    # encuentra los estilos de la plantilla de pandoc.
    estilos = {s.name: s for s in doc.styles}
    # Arial en todos los estilos (título, autor, índice, tablas…) y en los valores por defecto.
    for e in estilos.values():
        if e.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            arial(e.font, e.element.get_or_add_rPr())
    predeterminado = doc.styles.element.find(qn("w:docDefaults"))
    if predeterminado is not None:
        rpr = predeterminado.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            fuentes = rpr.find(qn("w:rFonts"))
            if fuentes is None:
                fuentes = OxmlElement("w:rFonts")
                rpr.insert(0, fuentes)
            quitar_tema(fuentes)
            for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fuentes.set(qn(atributo), "Arial")
    for nombre in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if nombre in estilos:
            e = estilos[nombre]
            fuente(e, 11)
            pf = e.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            pf.space_after = Pt(6)
            if nombre != "Compact":
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for nombre, tam in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11), ("Heading 4", 11)):
        e = estilos[nombre]
        fuente(e, tam, negrita=True)
        e.paragraph_format.space_before = Pt(12)
        e.paragraph_format.space_after = Pt(6)
    estilos["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    estilos["Heading 1"].paragraph_format.page_break_before = True  # Art. 35 h)
    for nombre in ("Caption", "Table Caption", "Image Caption", "Footnote Text"):
        if nombre in estilos:
            fuente(estilos[nombre], 10)
    if "Bibliography" in estilos:
        b = estilos["Bibliography"]
        fuente(b, 11)
        b.paragraph_format.left_indent = Cm(1.27)       # sangría francesa APA
        b.paragraph_format.first_line_indent = Cm(-1.27)
        b.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    # Bloques de código y diagramas de texto: fuente de ancho fijo, interlineado simple, sin justificar
    # (si no, los diagramas ASCII de los manuales se deforman).
    def estilo(nombre, tipo):
        return estilos.get(nombre) or doc.styles.add_style(nombre, tipo)

    codigo = estilo("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    codigo.font.name = "Liberation Mono"
    codigo.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Liberation Mono")
    codigo.font.size = Pt(8)
    pf = codigo.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = pf.space_after = Pt(0)
    caracter = estilo("Verbatim Char", WD_STYLE_TYPE.CHARACTER)
    caracter.font.name = "Liberation Mono"
    caracter.font.size = Pt(9)

    doc.save(DESTINO)
    tmp.unlink()
    print(f"Plantilla creada: {DESTINO}")


if __name__ == "__main__":
    main()
